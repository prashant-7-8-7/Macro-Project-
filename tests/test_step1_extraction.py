import sys
import tempfile
from pathlib import Path

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx
from src.ingestion.parsers import DocumentParser, normalize_document_text

def test_normalization():
    # 1. NFKC normalization
    # Fullwidth latin 'ＡＢＣ' -> 'ABC'
    raw = "Ｈｅｌｌｏ Ｗｏｒｌｄ"
    norm = normalize_document_text(raw)
    assert norm == "Hello World", f"Expected 'Hello World', got '{norm}'"

    # 2. Zero-width characters removal (\u200b, \ufeff, etc.)
    text_with_zw = "Hel\u200blo\uFEFF Wor\u200cld\u200d"
    norm_zw = normalize_document_text(text_with_zw)
    assert norm_zw == "Hello World", f"Zero-width chars not removed: '{norm_zw}'"

    # 3. Hyphenated line break removal
    text_hyphen = "This is multi-\nlingual and high-\r\nperformance text."
    norm_hyphen = normalize_document_text(text_hyphen)
    assert "multilingual" in norm_hyphen, f"Hyphen not joined: '{norm_hyphen}'"
    assert "highperformance" in norm_hyphen, f"Hyphen not joined: '{norm_hyphen}'"
    print("[OK] test_normalization passed")

def test_txt_encodings():
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        
        # UTF-8 with BOM (utf-8-sig)
        p_bom = tmp_path / "test_bom.txt"
        p_bom.write_bytes("हिंदी और Español".encode('utf-8-sig'))
        pages, warn = DocumentParser.parse(str(p_bom))
        assert len(pages) == 1
        assert "हिंदी" in pages[0]["text"]
        assert "Español" in pages[0]["text"]

        # UTF-16
        p_u16 = tmp_path / "test_u16.txt"
        p_u16.write_bytes("日本語とالعربية".encode('utf-16'))
        pages, warn = DocumentParser.parse(str(p_u16))
        assert len(pages) == 1
        assert "日本語" in pages[0]["text"]
        assert "العربية" in pages[0]["text"]

        # CP1252
        p_cp1252 = tmp_path / "test_cp1252.txt"
        p_cp1252.write_bytes("Café résumé".encode('cp1252'))
        pages, warn = DocumentParser.parse(str(p_cp1252))
        assert len(pages) == 1
        assert "Café" in pages[0]["text"]

        # Corrupt / bad bytes (must never crash)
        p_bad = tmp_path / "test_bad.txt"
        p_bad.write_bytes(b"\xff\xfe\x00\x80\x90\xa0\xbc\xde random text")
        pages, warn = DocumentParser.parse(str(p_bad))
        assert len(pages) == 1
        assert "random text" in pages[0]["text"]
        
    print("[OK] test_txt_encodings passed")

def test_docx_tables_and_scripts():
    with tempfile.TemporaryDirectory() as tmpdir:
        doc_path = Path(tmpdir) / "multilingual_tables.docx"
        doc = docx.Document()
        doc.add_paragraph("यह एक पैराग्राफ है (Hindi paragraph).")
        
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Spanish: Hola mundo"
        table.rows[0].cells[1].text = "Chinese: 你好世界"
        table.rows[1].cells[0].text = "Arabic: مرحبا بالعالم"
        table.rows[1].cells[1].text = "Russian: Привет мир"
        
        doc.save(str(doc_path))
        
        pages, warn = DocumentParser.parse(str(doc_path))
        assert len(pages) == 1
        text = pages[0]["text"]
        assert "यह एक पैराग्राफ है" in text
        assert "Hola mundo" in text
        assert "你好世界" in text
        assert "مرحبا بالعالم" in text
        assert "Привет мир" in text
        
    print("[OK] test_docx_tables_and_scripts passed")

def test_pdf_quality_check():
    # 1. Clean pages -> no warning
    clean_pages = [{"text": "Machine Learning and Neural Networks in Spanish and Hindi नमस्ते", "page_number": 1}]
    checked, warn = DocumentParser._check_pdf_quality(clean_pages)
    assert warn is None

    # 2. Empty pages -> warning
    empty_pages = []
    checked, warn = DocumentParser._check_pdf_quality(empty_pages)
    assert warn is not None
    assert "This PDF looks scanned or uses a non-standard font" in warn

    # 3. > 30% replacement characters -> warning
    garbage_pages = [{"text": "Hello \ufffd\ufffd\ufffd\ufffd\ufffd\ufffd\ufffd\ufffd", "page_number": 1}]
    checked, warn = DocumentParser._check_pdf_quality(garbage_pages)
    assert warn is not None
    assert "This PDF looks scanned or uses a non-standard font" in warn

    # 4. PUA characters -> warning
    pua_pages = [{"text": "\ue001\ue002\ue003\ue004\ue005 Normal", "page_number": 1}]
    checked, warn = DocumentParser._check_pdf_quality(pua_pages)
    assert warn is not None
    
    print("[OK] test_pdf_quality_check passed")

if __name__ == "__main__":
    test_normalization()
    test_txt_encodings()
    test_docx_tables_and_scripts()
    test_pdf_quality_check()
    print("\nALL STEP 1 EXTRACTION TESTS PASSED SUCCESSFULLY!")
