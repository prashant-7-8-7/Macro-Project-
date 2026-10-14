import os
import sys
from pathlib import Path
import docx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


def generate_all_fixtures():
    fixtures_dir = Path(__file__).resolve().parent / "fixtures"
    fixtures_dir.mkdir(parents=True, exist_ok=True)

    # 1. Hindi TXT
    hi_txt = (
        "कृत्रिम बुद्धिमत्ता और मशीन लर्निंग\n\n"
        "कृत्रिम बुद्धिमत्ता कंप्यूटर विज्ञान की वह शाखा है जो मशीनों को सोचने और निर्णय लेने की क्षमता प्रदान करती है।\n"
        "मशीन लर्निंग इसका एक महत्वपूर्ण उपक्षेत्र है जहाँ एल्गोरिदम डेटा से सीखकर पैटर्न पहचानते हैं।\n"
        "सुपरवाइज्ड लर्निंग में मॉडल्स को लेबल किए गए डेटा के साथ प्रशिक्षित किया जाता है ताकि वे सटीक भविष्यवाणियां कर सकें।\n"
        "अनसुपरवाइज्ड लर्निंग में डेटा बिना किसी लेबल के प्रदान किया जाता है और सिस्टम खुद छिपे हुए पैटर्न खोजता है।\n"
        "डीप लर्निंग न्यूरल नेटवर्क का उपयोग करके जटिल समस्याओं को हल करता है जैसे छवि पहचान और प्राकृतिक भाषा प्रसंस्करण।"
    )
    with open(fixtures_dir / "hindi.txt", "w", encoding="utf-8") as f:
        f.write(hi_txt)

    # 2. Spanish TXT
    es_txt = (
        "Energías Renovables y Sostenibilidad\n\n"
        "La energía solar es una de las fuentes renovables más limpias y abundantes del planeta.\n"
        "Los paneles fotovoltaicos convierten la radiación solar directamente en electricidad mediante el efecto fotoeléctrico.\n"
        "La energía eólica aprovecha la fuerza del viento para mover turbinas que generan electricidad de manera eficiente.\n"
        "La transición hacia fuentes renovables reduce drásticamente las emisiones de gases de efecto invernadero.\n"
        "El almacenamiento mediante baterías de iones de litio permite garantizar el suministro constante de energía durante la noche."
    )
    with open(fixtures_dir / "spanish.txt", "w", encoding="utf-8") as f:
        f.write(es_txt)

    # 3. Arabic TXT
    ar_txt = (
        "الذكاء الاصطناعي ومعالجة اللغات الطبيعية\n\n"
        "يعتبر الذكاء الاصطناعي أحد أهم التطورات التكنولوجية في العصر الحديث.\n"
        "معالجة اللغات الطبيعية هي فرع من فروع الذكاء الاصطناعي يهدف إلى تمكين الحواسيب من فهم النصوص البشرية وتوليدها.\n"
        "تعتمد النماذج اللغوية الحديثة على معمارية المحولات لتحليل السياق بدقة عالية عبر مختلف اللغات.\n"
        "تطبيقات الذكاء الاصطناعي تشمل الترجمة الآلية، والتحليل الدلالي، والمساعدات الصوتية الذكية.\n"
        "تسهم هذه التقنيات في تحسين الإنتاجية وتقديم حلول مبتكرة في مختلف المجالات الصناعية والطبية والتعليمية."
    )
    with open(fixtures_dir / "arabic.txt", "w", encoding="utf-8") as f:
        f.write(ar_txt)

    # 4. Chinese TXT
    zh_txt = (
        "自然语言处理与大语言模型\n\n"
        "自然语言处理是人工智能领域的重要研究方向，致力于让计算机理解和生成人类语言。\n"
        "大语言模型基于深度神经网络，通过海量文本数据的预训练学习丰富的语言和世界常识。\n"
        "词嵌入技术将文本转换为高维向量空间中的连续表示，准确捕捉语义之间的关联与距离。\n"
        "注意力机制使得模型能够动态关注输入序列中不同位置的关键信息，解决长距离依赖难题。\n"
        "这些技术广泛应用于机器翻译、信息检索、智能问答以及自动化内容生成等前沿场景。"
    )
    with open(fixtures_dir / "chinese.txt", "w", encoding="utf-8") as f:
        f.write(zh_txt)

    # 5. English TXT (Regression fixture)
    en_txt = (
        "Cloud Computing Architecture and Scalability\n\n"
        "Cloud computing provides on-demand availability of computer system resources such as data storage and computing power.\n"
        "Infrastructure as a Service (IaaS) offers fundamental computing infrastructure like virtual machines and raw block storage.\n"
        "Platform as a Service (PaaS) delivers hardware and software tools over the internet for agile application development.\n"
        "Software as a Service (SaaS) provides complete software applications accessible directly via modern web browsers.\n"
        "Cloud elasticity allows organizations to scale computing capacity dynamically according to current real-time workload."
    )
    with open(fixtures_dir / "english.txt", "w", encoding="utf-8") as f:
        f.write(en_txt)

    # 6. Multilingual DOCX with Tables & Non-Latin scripts
    doc = docx.Document()
    doc.add_heading("Multilingual Document Intelligence", level=1)
    doc.add_paragraph("This document contains multilingual text and tabular information across various scripts.")
    doc.add_paragraph("Spanish: La tecnología moderna impulsa la transformación digital en todo el mundo.")
    doc.add_paragraph("Hindi: आधुनिक तकनीक दुनिया भर में डिजिटल परिवर्तन को तीव्र गति दे रही है।")

    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Language"
    hdr_cells[1].text = "Concept"
    hdr_cells[2].text = "Description"

    row1 = table.add_row().cells
    row1[0].text = "Hindi"
    row1[1].text = "मशीन लर्निंग"
    row1[2].text = "कंप्यूटर एल्गोरिदम जो डेटा और पूर्व अनुभवों से सीखते हैं।"

    row2 = table.add_row().cells
    row2[0].text = "Spanish"
    row2[1].text = "Inteligencia Artificial"
    row2[2].text = "Sistemas capaces de realizar tareas complejas de razonamiento y predicción."

    doc.save(str(fixtures_dir / "multilingual_sample.docx"))

    # 7. Spanish PDF
    c_es = canvas.Canvas(str(fixtures_dir / "spanish.pdf"), pagesize=letter)
    c_es.drawString(100, 750, "Energias Renovables y Sostenibilidad")
    c_es.drawString(100, 720, "La energia solar es una de las fuentes renovables mas limpias del planeta.")
    c_es.drawString(100, 690, "Los paneles fotovoltaicos convierten la radiacion solar en electricidad.")
    c_es.drawString(100, 660, "La transicion energetica reduce drásticamente las emisiones contaminantes.")
    c_es.save()

    # 8. Hindi PDF
    # Note: Hindi text requires a TrueType Unicode font (e.g. Nirmala.ttc or Mangal.ttf on Windows)
    # Default PDF standard fonts (Helvetica, Times) do not include Devanagari glyphs.
    nirmala_path = "C:/Windows/Fonts/Nirmala.ttc"
    if os.path.exists(nirmala_path):
        try:
            pdfmetrics.registerFont(TTFont("Nirmala", nirmala_path, subfontIndex=0))
            c_hi = canvas.Canvas(str(fixtures_dir / "hindi.pdf"), pagesize=letter)
            c_hi.setFont("Nirmala", 14)
            c_hi.drawString(100, 750, "कृत्रिम बुद्धिमत्ता और मशीन लर्निंग")
            c_hi.drawString(100, 720, "कृत्रिम बुद्धिमत्ता कंप्यूटर विज्ञान की एक महत्वपूर्ण शाखा है।")
            c_hi.drawString(100, 690, "मशीन लर्निंग में एल्गोरिदम डेटा से सीखकर सटीक निर्णय लेते हैं।")
            c_hi.drawString(100, 660, "सुपरवाइज्ड लर्निंग में मॉडल्स को लेबल किए गए डेटा से सिखाया जाता है।")
            c_hi.save()
            print("[NOTE] Hindi PDF created successfully with Windows Unicode font: Nirmala.ttc")
        except Exception as e:
            print(f"[NOTE] Could not register Nirmala font ({e}); falling back to standard font for Hindi PDF.")
            _create_fallback_hindi_pdf(fixtures_dir / "hindi.pdf")
    else:
        print("[NOTE] Nirmala.ttc not found; creating fallback Hindi PDF.")
        _create_fallback_hindi_pdf(fixtures_dir / "hindi.pdf")

    print("[OK] All fixtures generated successfully in tests/fixtures/")


def _create_fallback_hindi_pdf(target_path: Path):
    c = canvas.Canvas(str(target_path), pagesize=letter)
    c.drawString(100, 750, "कृत्रिम बुद्धिमत्ता और मशीन लर्निंग")
    c.drawString(100, 720, "मशीन लर्निंग में एल्गोरिदम डेटा से सीखते हैं।")
    c.save()


if __name__ == "__main__":
    generate_all_fixtures()
