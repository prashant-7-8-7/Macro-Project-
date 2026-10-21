import csv
import json
import re
import unicodedata
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional


def normalize_document_text(text: str) -> str:
    """
    Normalizes text using Unicode NFKC normalization, strips zero-width characters,
    and removes hyphenated line breaks.
    """
    if not text:
        return ""
    # 1. NFKC normalization
    text = unicodedata.normalize("NFKC", text)
    # 2. Strip zero-width characters (ZWSP, ZWNJ, ZWJ, BOM, LTR/RTL marks)
    text = re.sub(r'[\u200B-\u200F\uFEFF]', '', text)
    # 3. Remove hyphenated line breaks (e.g., "informa-\ntion" -> "information")
    text = re.sub(r'(\w+)-\s*[\r\n]+\s*(\w+)', r'\1\2', text, flags=re.UNICODE)
    return text


class DocumentParser:
    @staticmethod
    def _decode_bytes(raw_bytes: bytes) -> str:
        """
        Detects encoding and decodes raw bytes without crashing on malformed bytes:
        tries utf-8-sig, utf-16, charset-normalizer, cp1252, and finally utf-8 with replacement.
        """
        # 1. Try utf-8-sig (handles UTF-8 with or without BOM)
        try:
            return raw_bytes.decode('utf-8-sig')
        except UnicodeDecodeError:
            pass

        # 2. Try utf-16 if BOM present or clear UTF-16 patterns
        if raw_bytes.startswith((b'\xff\xfe', b'\xfe\xff')) or (len(raw_bytes) >= 4 and b'\x00' in raw_bytes[:16]):
            try:
                return raw_bytes.decode('utf-16')
            except UnicodeDecodeError:
                pass

        # 3. Try charset-normalizer detection
        detected_enc = None
        try:
            import charset_normalizer
            matches = charset_normalizer.from_bytes(raw_bytes)
            best = matches.best()
            if best and best.encoding:
                norm_enc = best.encoding.lower().replace('-', '').replace('_', '')
                if best.coherence > 0 or norm_enc in ('utf8', 'utf16', 'gbk', 'gb2312', 'gb18030', 'big5', 'shiftjis', 'eucjp', 'euckr', 'cp932', 'cp936', 'cp949', 'cp950'):
                    detected_enc = best.encoding
        except Exception:
            pass

        if detected_enc:
            try:
                return raw_bytes.decode(detected_enc)
            except (UnicodeDecodeError, LookupError):
                pass

        # 4. Try cp1252 (the standard ANSI fallback for single-byte text)
        try:
            return raw_bytes.decode('cp1252')
        except (UnicodeDecodeError, LookupError):
            pass

        # If charset-normalizer had any guess, try it before fallback
        try:
            import charset_normalizer
            res = charset_normalizer.detect(raw_bytes)
            if res and res.get('encoding'):
                return raw_bytes.decode(res['encoding'], errors='replace')
        except Exception:
            pass

        # 5. Fallback: Never crash on bad bytes
        return raw_bytes.decode('utf-8', errors='replace')

    @staticmethod
    def _check_pdf_quality(pages: List[Dict]) -> Tuple[List[Dict], Optional[str]]:
        warning_msg = "This PDF looks scanned or uses a non-standard font, so text could not be read reliably"
        if not pages:
            return [], warning_msg

        full_text = "".join(p.get("text", "") for p in pages)
        non_ws_chars = [c for c in full_text if not c.isspace()]

        if not non_ws_chars:
            return pages, warning_msg

        garbage_count = 0
        for c in non_ws_chars:
            code = ord(c)
            cat = unicodedata.category(c)
            # Replacement character \ufffd, null byte \x00, Private Use Area, or unassigned/control
            if (
                c == '\ufffd'
                or c == '\x00'
                or (0xE000 <= code <= 0xF8FF)
                or (0xF0000 <= code <= 0xFFFFD)
                or (0x100000 <= code <= 0x10FFFD)
                or cat in ('Cc', 'Co', 'Cs', 'Cn')
            ):
                garbage_count += 1

        garbage_ratio = garbage_count / len(non_ws_chars)
        if garbage_ratio > 0.30:
            return pages, warning_msg

        return pages, None

    @staticmethod
    def parse(file_path: str) -> Tuple[List[Dict[str, Any]], Optional[str]]:
        """
        Parses a file and returns a tuple: (pages, warning_message).
        Each page item is dict: {'text': str, 'page_number': int}
        warning_message is an optional warning string or None.
        """
        path = Path(file_path)
        ext = path.suffix.lower()

        if ext == '.pdf':
            return DocumentParser._parse_pdf(path)
        elif ext in ['.txt', '.md']:
            return DocumentParser._parse_text(path)
        elif ext == '.docx':
            return DocumentParser._parse_docx(path)
        elif ext == '.csv':
            return DocumentParser._parse_csv(path)
        elif ext == '.json':
            return DocumentParser._parse_json(path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    @staticmethod
    def _parse_pdf(path: Path) -> Tuple[List[Dict], Optional[str]]:
        import os
        import fitz
        from src.ingestion.ocr import ocr_page
        pages, warnings, ocr_count = [], set(), 0
        with fitz.open(path) as document:
            if document.needs_pass:
                raise ValueError("Password-protected PDFs are not supported")
            for number, page in enumerate(document, 1):
                text = normalize_document_text(page.get_text()).strip()
                words = []
                _, warning = DocumentParser._check_pdf_quality([{'text': text}])
                if len(text) < 40 or warning:
                    if ocr_count < int(os.environ.get('OCR_MAX_PAGES', '10')):
                        ocr_count += 1
                        try:
                            extracted, words, ocr_warning = ocr_page(page)
                            if extracted.strip():
                                text = normalize_document_text(extracted).strip()
                            if ocr_warning:
                                warnings.add(ocr_warning)
                        except Exception:
                            warnings.add(f'OCR could not read page {number}. Check local OCR languages and image quality.')
                    else:
                        warnings.add('OCR page limit reached; increase OCR_MAX_PAGES for larger scans.')
                if text:
                    pages.append({'text': text, 'page_number': number, 'ocr_words': words})
                else:
                    warnings.add(f'No readable text on page {number}.')
        if not pages:
            warnings.add('This PDF looks scanned or uses a non-standard font, so text could not be read reliably')
        return pages, ' '.join(sorted(warnings)) or None

    @staticmethod
    def _parse_text(path: Path) -> Tuple[List[Dict], Optional[str]]:
        raw_bytes = path.read_bytes()
        text = DocumentParser._decode_bytes(raw_bytes)
        text = normalize_document_text(text).strip()
        if not text:
            return [], None
        return [{"text": text, "page_number": 1}], None

    @staticmethod
    def _parse_docx(path: Path) -> Tuple[List[Dict], Optional[str]]:
        try:
            from docx import Document
            from docx.text.paragraph import Paragraph
            from docx.table import Table

            doc = Document(str(path))
            elements_text = []

            # Extract paragraphs and tables preserving visual reading order
            for element in doc.element.body:
                if element.tag.endswith('p'):
                    p = Paragraph(element, doc)
                    if p.text and p.text.strip():
                        elements_text.append(p.text.strip())
                elif element.tag.endswith('tbl'):
                    tbl = Table(element, doc)
                    for row in tbl.rows:
                        seen_cells = set()
                        cell_texts = []
                        for cell in row.cells:
                            if cell._tc not in seen_cells:
                                seen_cells.add(cell._tc)
                                c_text = cell.text.strip()
                                if c_text:
                                    cell_texts.append(c_text)
                        if cell_texts:
                            elements_text.append(" | ".join(cell_texts))

            # Fallback if doc.element.body iteration produced nothing
            if not elements_text:
                for p in doc.paragraphs:
                    if p.text and p.text.strip():
                        elements_text.append(p.text.strip())
                for tbl in doc.tables:
                    for row in tbl.rows:
                        seen_cells = set()
                        cell_texts = []
                        for cell in row.cells:
                            if cell._tc not in seen_cells:
                                seen_cells.add(cell._tc)
                                c_text = cell.text.strip()
                                if c_text:
                                    cell_texts.append(c_text)
                        if cell_texts:
                            elements_text.append(" | ".join(cell_texts))

            full_text = "\n\n".join(elements_text)
            full_text = normalize_document_text(full_text).strip()
            if not full_text:
                return [], None
            return [{"text": full_text, "page_number": 1}], None
        except Exception as e:
            raise RuntimeError(f"Error parsing DOCX: {e}")

    @staticmethod
    def _parse_csv(path: Path) -> Tuple[List[Dict], Optional[str]]:
        raw_bytes = path.read_bytes()
        decoded_text = DocumentParser._decode_bytes(raw_bytes)
        lines = []
        reader = csv.reader(decoded_text.splitlines())
        for row in reader:
            if row:
                lines.append(" | ".join(row))
        full_text = "\n".join(lines)
        full_text = normalize_document_text(full_text).strip()
        if not full_text:
            return [], None
        return [{"text": full_text, "page_number": 1}], None

    @staticmethod
    def _parse_json(path: Path) -> Tuple[List[Dict], Optional[str]]:
        raw_bytes = path.read_bytes()
        decoded_text = DocumentParser._decode_bytes(raw_bytes)
        data = json.loads(decoded_text)
        formatted_text = json.dumps(data, indent=2, ensure_ascii=False)
        formatted_text = normalize_document_text(formatted_text).strip()
        if not formatted_text:
            return [], None
        return [{"text": formatted_text, "page_number": 1}], None
